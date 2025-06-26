# chatbot_api.py

from fastapi import FastAPI, HTTPException, Depends
from pydantic import BaseModel
import uvicorn
import google.generativeai as genai
import os
from dotenv import load_dotenv
import pandas as pd
from datetime import datetime
from docx import Document
# from fastapi.staticfiles import StaticFiles # 정적 파일 서빙을 하지 않으므로 주석 처리 또는 제거
from fastapi.responses import HTMLResponse, FileResponse
from fastapi.middleware.cors import CORSMiddleware
import re

# SQLAlchemy 관련 임포트
from sqlalchemy import create_engine, Column, Integer, String, Text, DateTime, ForeignKey
from sqlalchemy.ext.declarative import declarative_base
from sqlalchemy.orm import sessionmaker, Session
from sqlalchemy.orm.exc import NoResultFound
from sqlalchemy import func # <<< 이제 이 줄은 완벽하게 있습니다! 잘하셨습니다.

# --- FastAPI 앱 인스턴스 생성 ---
app = FastAPI()

# --- CORS 미들웨어 추가 ---
origins = ["*"] # 개발용: 모든 Origin 허용
app.add_middleware(
    CORSMiddleware,
    allow_origins=origins,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# --- 정적 파일 서비스 설정 (제거됨 - FastAPI는 API만 제공) ---
# 이 부분은 FastAPI가 프론트엔드 파일을 직접 서빙할 때 필요합니다.
# 사용자님의 프로젝트 구조에서는 다른 웹 서버가 이 역할을 하므로 제거합니다.

# --- chatbot.html 파일을 직접 제공하는 엔드포인트 (제거됨 - FastAPI는 API만 제공) ---
@app.get("/", response_class=HTMLResponse)
async def read_root_html_placeholder():
    raise HTTPException(status_code=404, detail="FastAPI is an API server. Access frontend via your main web server (e.g., http://localhost:8080/chatbot.html).")

@app.get("/chatbot.html", response_class=HTMLResponse)
async def get_chatbot_html_placeholder():
    raise HTTPException(status_code=404, detail="FastAPI is an API server. Access frontend via your main web server (e.g., http://localhost:8080/chatbot.html).")


# --- 데이터베이스 설정 ---
DATABASE_URL = "mysql+pymysql://root:user1234@192.168.0.30:3306/greenhand" 

engine = create_engine(DATABASE_URL) 
SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
Base = declarative_base()

# --- DB 모델 정의 ---
class User(Base):
    __tablename__ = "users"
    user_id = Column(Integer, primary_key=True, index=True)
    username = Column(String(50), unique=True, index=True, nullable=False)
    password_hash = Column(String(500), nullable=False, default='')
    nickname = Column(String(50), nullable=False, default='익명')
    name = Column(String(50), nullable=False, default='이름없음')
    email = Column(String(100), unique=True, index=True, nullable=True)
    address = Column(String(255), nullable=True)
    phone_number = Column(String(20), nullable=True)
    created_at = Column(DateTime, nullable=False, default=datetime.now)
    last_login_at = Column(DateTime, nullable=True)

    def __repr__(self):
        return f"<User(user_id={self.user_id}, username='{self.username}', nickname='{self.nickname}')>"

class ChatLog(Base):
    __tablename__ = "chat_logs"
    id = Column(Integer, primary_key=True, autoincrement=True)
    user_id = Column(Integer, ForeignKey("users.user_id"), nullable=False)
    message_type = Column(String(10), nullable=False)
    message_content = Column(Text, nullable=False)
    timestamp = Column(DateTime, default=datetime.now)

    def __repr__(self):
        return f"<ChatLog(id={self.id}, user_id={self.user_id}, message_type='{self.message_type}', timestamp='{self.timestamp}')>"


# --- API 키 및 모델 초기화 ---
load_dotenv(dotenv_path='test.env')
GOOGLE_API_KEY = os.getenv("GOOGLE_API_KEY")

if not GOOGLE_API_KEY:
    print("Google API 키가 설정되지 않았습니다. test.env 파일을 확인해주세요.")
    raise ValueError("GOOGLE_API_KEY 환경 변수가 설정되지 않았습니다. test.env 파일을 확인해주세요.")

genai.configure(api_key=GOOGLE_API_KEY)

try:
    model = genai.GenerativeModel('gemini-1.5-flash')
except Exception as e:
    print(f"모델 초기화 실패: {e}")
    raise Exception(f"Gemini 모델 초기화에 실패했습니다. API 키, 결제, 모델 가용성을 확인하세요: {e}")

# --- 지식 데이터베이스 (농업 정보) ---
agricultural_knowledge_base = {
    "상추 재배 방법": "상추는 서늘하고 햇볕이 잘 드는 곳에서 잘 자랍니다. 씨앗을 심고 싹이 나면 솎아주세요. 물은 흙이 마르지 않게 꾸준히 주는 것이 중요합니다.",
    "토마토 병충해": "토마토에 흔한 병충해로는 탄저병, 역병, 온실가루이 등이 있습니다. 각 병충해에 맞는 방제법을 사용해야 합니다.",
    "딸기 수확 시기": "딸기는 보통 4월에서 6월 사이에 수확하며, 품종과 재배 환경에 따라 달라질 수 있습니다.",
    "스마트팜이란?": "스마트팜은 정보통신기술(ICT)을 활용하여 작물 생육 환경을 원격 및 자동으로 제어하는 농업 시스템입니다.",
    "온도 조절 중요성": "작물 생육에 있어 온도는 매우 중요합니다. 너무 높거나 낮은 온도는 작물의 스트레스를 유발하고 성장을 저해할 수 있습니다.",
    "질병 진단 기능": "식물 잎사귀 이미지를 업로드하면 질병 또는 해충을 진단해 드립니다. [질병 진단 바로가기](/PRH-002)",
    "시뮬레이션 기능": "재배할 농작물 종류를 선택하고 환경 조건(온도, 습도, 일조량 등)을 입력하면 미래 생육 상태(예상 수확량, 성장 곡선)를 예측하여 시각화해야 한다. [농작물 시뮬레이션 바로가기](/PRH-003)",
    "내 농장 확인": "현재 키우시는 작물의 성장 진행도를 한 눈에 확인하고 관리할 수 있습니다. [내 농장 바로가기](/my_farm_page_id)",
    "엑셀 보고서": "엑셀 보고서 생성을 요청하셨습니다. 'generate_excel_report' 함수를 호출합니다.",
    "워드 보고서": "워드 보고서 생성을 요청하셨습니다. 'generate_word_report' 함수를 호출합니다."
}

# --- 보고서/엑셀 생성 함수 ---
def generate_excel_report(query_text):
    try:
        data = {
            '날짜': [f'2025-06-{i:02d}' for i in range(10, 17)],
            '상추_성장(cm)': [5, 6, 7.5, 8, 9, 9.5, 10],
            '온도(℃)': [22, 23, 21, 24, 23, 22, 25],
            '습도(%)': [60, 62, 58, 65, 63, 61, 64]
        }
        df = pd.DataFrame(data)

        # 보고서 파일은 FastAPI 서버의 실행 디렉토리에 'reports' 폴더 안에 생성됩니다.
        file_name = f"reports/상추_성장_보고서_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        os.makedirs(os.path.dirname(file_name), exist_ok=True)
        df.to_excel(file_name, index=False)
        return {"message": f"요청하신 엑셀 보고서 '{file_name}'가 성공적으로 생성되었습니다.", "file_path": file_name}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"엑셀 보고서 생성 중 예기치 않은 오류 발생: {e}")

def generate_word_report(query_text):
    try:
        document = Document()
        document.add_heading('농작물 성장 보고서', level=1)
        document.add_paragraph(f'생성 날짜: {datetime.now().strftime("%Y년 %m월 %d일")}')
        document.add_heading('1. 상추 성장 개요', level=2)
        document.add_paragraph('이 보고서는 특정 기간 동안의 상추 성장 데이터를 요약합니다. 데이터는 센서 모니터링 시스템에서 수집되었습니다.')
        document.add_heading('2. 주요 데이터', level=2)
        data = {
            '날짜': [f'2025-06-{i:02d}' for i in range(10, 17)],
            '상추_성장(cm)': [5, 6, 7.5, 8, 9, 9.5, 10],
            '온도(℃)': [22, 23, 21, 24, 23, 22, 25],
            '습度(%)': [60, 62, 58, 65, 63, 61, 64]
        }
        df = pd.DataFrame(data)
        document.add_paragraph(df.to_string())
        # 보고서 파일은 FastAPI 서버의 실행 디렉토리에 'reports' 폴더 안에 생성됩니다.
        file_name = f"reports/농작물_성장_보고서_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
        os.makedirs(os.path.dirname(file_name), exist_ok=True)
        document.save(file_name)
        return {"message": f"요청하신 워드 보고서 '{file_name}'가 성공적으로 생성되었습니다.", "file_path": file_name}
    except ImportError:
        raise HTTPException(status_code=500, detail="docx 라이브러리가 설치되지 않았습니다. 'pip install python-docx'를 실행하여 설치해주세요.")
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"워드 보고서 생성 중 예기치 않은 오류 발생: {e}")

# --- DB 세션을 위한 의존성 주입 함수 ---
def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()

# --- 사용자 조회 또는 생성 함수 ---
def get_or_create_user(db: Session, user_id: int, username: str = "default_user"):
    try:
        user = db.query(User).filter(User.user_id == user_id).one()
        print(f"Existing user found: {user}")
        return user
    except NoResultFound:
        print(f"User with user_id={user_id} not found, creating new user.")
        new_user = User(
            user_id=user_id, 
            username=username,
            password_hash='',
            nickname='익명사용자',
            name='기본사용자',
            created_at=datetime.now()
        ) 
        db.add(new_user)
        db.commit()
        db.refresh(new_user)
        print(f"New user created: {new_user}")
        return new_user
    except Exception as e:
        print(f"Error getting or creating user: {e}")
        db.rollback()
        raise

# --- 챗봇 응답 생성 함수 (핵심 로직 및 로그 저장) ---
def get_chatbot_response(user_query: str, db: Session, user_id: int = 1):
    print(f"User query received by FastAPI: {user_query}")

    try:
        user = get_or_create_user(db, user_id)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"사용자 정보를 처리하는 중 오류가 발생했습니다: {e}")

    try:
        user_log = ChatLog(user_id=user_id, message_type="user", message_content=user_query, timestamp=datetime.now())
        db.add(user_log)
        db.commit()
        db.refresh(user_log)
    except Exception as e:
        db.rollback()
        print(f"Failed to log user message: {e}")
        raise HTTPException(status_code=500, detail=f"사용자 메시지 기록 중 오류 발생: {e}")

    prompt = f"""
    당신은 스마트 농업 도우미 챗봇입니다. 사용자의 질문에 대해 친절하게 답변하고,
    필요 시 다른 기능(질병 진단, 시뮬레이션, 내 농장)으로의 **Markdown 링크**를 포함하여 안내합니다.
    링크는 반드시 `[텍스트](링크주소)` 형태로 생성해 주세요. 예를 들어 `[질병 진단 바로가기](/PRH-002)` 처럼요.
    링크주소는 보고서의 화면 ID(/PRH-002, /PRH-003, /my_farm_page_id)를 사용해주세요.
    만약 보고서나 엑셀 파일 생성을 요청하는 경우, 해당 기능을 수행해야 합니다.

    [농업 정보 제공 예시]
    질문: 상추는 어떻게 키우나요?
    답변: 상추는 서늘하고 햇볕이 잘 드는 곳에서 잘 자랍니다. 씨앗을 심고 싹이 나면 솎아주세요. 물은 흙이 마르지 않게 꾸준히 주는 것이 중요합니다.

    [기능 링크 제공 예시]
    질문: 상추가 시들시들해요. 왜 그런가요?
    답변: 상추의 상태를 정확히 진단하려면 질병 진단 기능을 사용해 보세요. [질병 진단 바로가기](/PRH-002)
     질문: 상추를 키우면 수확량이 얼마나 될까요?
     답변: 예상 수확량을 확인하려면 농작물 시뮬레이션 기능을 이용해 보세요. [농작물 시뮬레이션 바로가기](/PRH-003)
    질문: 제가 키우고 있는 상추 상태를 보고 싶어요.
    답변: 현재 키우시는 작물의 성장 진행도는 '내 농장' 페이지에서 확인하실 수 있습니다. [내 농장 바로가기](/my_farm_page_id)

    [보고서/엑셀 생성 예시]
    질문: 지난주 상추 성장 데이터를 엑셀로 정리해줘.
    답변: 엑셀 파일 생성을 요청하셨습니다.
    질문: 이번 달 농작물 보고서를 워드 파일로 만들어줘.
    답변: 보고서 생성을 요청하셨습니다.

    사용자 질문: {user_query}
    답변:
    """

    try:
        response_content = ""

        if ("엑셀" in user_query or "excel" in user_query) and ("생성" in user_query or "만들어" in user_query or "정리" in user_query):
            file_result = generate_excel_report(user_query)
            bot_log = ChatLog(user_id=user_id, message_type="bot", message_content=file_result["message"], timestamp=datetime.now())
            db.add(bot_log)
            db.commit()
            db.refresh(bot_log)
            return {"response": file_result["message"], "file_path": file_result["file_path"]}
        elif ("보고서" in user_query or "report" in user_query) and ("생성" in user_query or "만들어" in user_query or "워드" in user_query or "word" in user_query):
            file_result = generate_word_report(user_query)
            bot_log = ChatLog(user_id=user_id, message_type="bot", message_content=file_result["message"], timestamp=datetime.now())
            db.add(bot_log)
            db.commit()
            db.refresh(bot_log)
            return {"response": file_result["message"], "file_path": file_result["file_path"]}
        
        if "질병 진단" in user_query or "시들" in user_query or "병충해" in user_query or "잎사귀 이미지" in user_query:
            response_content = f"네, 식물 잎사귀 이미지를 통해 질병 진단을 도와드릴 수 있습니다."
            llm_link_id = "/PRH-002"
            response_content += f" [질병 진단 바로가기]({llm_link_id})"
        elif "시뮬레이션" in user_query or "수확량" in user_query or "성장 예측" in user_query or "환경 조건" in user_query:
            response_content = f"재배할 농작물 종류와 환경 조건을 입력하시면 미래 생육 상태를 예측해 드립니다."
            llm_link_id = "/PRH-003"
            response_content += f" [농작물 시뮬레이션 바로가기]({llm_link_id})"
        elif "내 농장" in user_query or "얼마나 컸는지" in user_query or "상태 확인" in user_query or "성장 진행도" in user_query:
            response_content = f"현재 키우시는 작물의 성장 진행도는 '내 농장' 페이지에서 확인하실 수 있습니다."
            llm_link_id = "/my_farm_page_id"
            response_content += f" [내 농장 바로가기]({llm_link_id})"
        else: 
            found_in_knowledge_base = False
            for keyword, answer in agricultural_knowledge_base.items():
                if keyword not in ["엑셀 보고서", "워드 보고서"] and keyword in user_query:
                    response_content = answer
                    found_in_knowledge_base = True
                    break
            
            if not found_in_knowledge_base:
                llm_response = model.generate_content(prompt)
                generated_text = llm_response.text
                response_content = generated_text.strip()
        
        markdown_link_pattern = re.compile(r'\[(.*?)\]\((.*?)\)')

        def replace_link_with_html(match):
            link_text = match.group(1)
            link_id = match.group(2)

            mapped_file = ""
            if link_id == "/PRH-002":
                mapped_file = "4.html"
            elif link_id == "/PRH-003":
                mapped_file = "2.html"
            elif link_id == "/my_farm_page_id":
                mapped_file = "main.html"
            else:
                mapped_file = "main.html"

            return f'<a href="{mapped_file}" class="link-button" onclick="event.preventDefault(); window.location.href=\'{mapped_file}\'">{link_text}</a>'

        processed_response_content = markdown_link_pattern.sub(replace_link_with_html, response_content)

        try:
            bot_log = ChatLog(user_id=user_id, message_type="bot", message_content=processed_response_content, timestamp=datetime.now())
            db.add(bot_log)
            db.commit()
            db.refresh(bot_log)
        except Exception as e:
            db.rollback()
            print(f"Failed to log bot message: {e}")
            raise HTTPException(status_code=500, detail=f"챗봇 응답 기록 중 오류 발생: {e}")

        return {"response": processed_response_content}

    except Exception as e:
        print(f"Unhandled error in get_chatbot_response: {e}")
        if "404 models/gemini-1.5-flash is not found" in str(e):
            raise HTTPException(status_code=500, detail="모델을 찾을 수 없거나 접근 권한이 없습니다. API 키, 결제, 모델 가용성을 확인하세요.")
        else:
            raise HTTPException(status_code=500, detail=f"챗봇 응답 생성 중 예기치 않은 오류 발생: {e}")

class ChatRequest(BaseModel):
    user_query: str
    user_id: int = 1

@app.post("/chatbot/ask")
async def ask_chatbot(request: ChatRequest, db: Session = Depends(get_db)):
    response_data = get_chatbot_response(request.user_query, db, request.user_id)
    return response_data

@app.get("/files/{file_path:path}")
async def download_file(file_path: str):
    if not file_path.startswith("reports/"):
        raise HTTPException(status_code=403, detail="Forbidden file path")
    
    full_path = os.path.join(os.getcwd(), file_path)

    if not os.path.exists(full_path):
        raise HTTPException(status_code=404, detail="File not found")
    
    if full_path.endswith(".xlsx"):
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    elif full_path.endswith(".docx"):
        media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    else:
        media_type = "application/octet-stream"

    return FileResponse(path=full_path, media_type=media_type, filename=os.path.basename(full_path))

# --- 새로운 엔드포인트: 이전 질문 내역 가져오기 ---
@app.get("/chatbot/history/{user_id}", response_model=list[str])
async def get_chat_history(user_id: int, db: Session = Depends(get_db)):
    """특정 사용자의 이전 질문 내역을 가져옵니다 (중복 제거, 최신 순)."""
    try:
        distinct_messages = db.query(ChatLog.message_content)\
                            .filter(ChatLog.user_id == user_id, ChatLog.message_type == "user")\
                            .group_by(ChatLog.message_content)\
                            .order_by(func.max(ChatLog.timestamp).desc())\
                            .all()
        return [msg[0] for msg in distinct_messages]
    except Exception as e:
        print(f"Error fetching chat history for user_id {user_id}: {e}")
        raise HTTPException(status_code=500, detail=f"이전 질문 내역을 불러오는 중 오류가 발생했습니다: {e}")


if __name__ == "__main__":
    os.makedirs("reports", exist_ok=True)
    
    Base.metadata.create_all(bind=engine) 
    print("Database tables created/checked.")

    db = SessionLocal()
    try:
        get_or_create_user(db, user_id=1, username="default_user")
    finally:
        db.close()
    
    uvicorn.run(app, host="0.0.0.0", port=8000)
